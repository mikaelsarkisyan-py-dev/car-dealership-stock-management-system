import openpyxl
import pytest

from stock import (Car, CarFilter, StockDataError, StockRepository, ValidationError, brand_counts,
                   build_car_fields, cars_for, distinct_values, filter_cars, model_counts,
                   price_distribution)


@pytest.fixture
def cars(stock_repo):
    return stock_repo.all()


def test_all_rows_loaded(cars):
    assert len(cars) == 36
    assert len({c.stock_id for c in cars}) == 36


def test_na_values_become_none(cars):
    tourneo = next(c for c in cars if c.model == "Tourneo connect MPV")
    assert tourneo.top_speed is None and tourneo.acceleration is None
    assert "N/A" in tourneo.spec_rows()[1][1]


def test_search_is_case_insensitive_substring(cars):
    assert {c.model for c in filter_cars(cars, CarFilter(text="gt86"))} == {"GT86"}
    assert len(filter_cars(cars, CarFilter(text="bmw"))) == 5
    assert len(filter_cars(cars, CarFilter(text="porsche 911"))) == 2


def test_filters_combine_with_and(cars):
    result = filter_cars(cars, CarFilter(make="ford", condition="New", transmission="Manual"))
    assert {c.model for c in result} == {"Puma", "Fiesta Active"}


def test_price_range_is_inclusive(cars):
    result = filter_cars(cars, CarFilter(min_price=900, max_price=1000))
    assert sorted(c.price for c in result) == [900, 1000]


def test_distinct_values_ignores_case(cars):
    assert distinct_values(cars, "colour").count("Blue") == 1


def test_brand_counts_include_zero_stock_brands(cars):
    counts = brand_counts(cars, known_brands={"Lotus"})
    assert counts["Lotus"] == 0 and counts["Ford"] == 6
    assert list(counts) == sorted(counts, key=str.casefold)


def test_model_counts_and_cars_for(cars):
    assert model_counts(cars, "Porsche") == {"911 Coupe": 1, "911 GT3": 1, "Cayenne Coupe": 1}
    assert len(cars_for(cars, "ford", "puma")) == 1


def test_price_bands():
    def car(price):
        return Car(1, "X", "Y", "New", price, 1, None, None, "", "", "", "", None, None,
                   "N/A", "", 2020, 0)
    assert price_distribution([]) == []
    result = price_distribution([car(12_250), car(14_999), car(21_000)])
    assert result == [("£10k - £15k", 2), ("£15k - £20k", 0), ("£20k - £25k", 1)]
    assert price_distribution([car(15_000)]) == [("£15k - £20k", 1)]  # boundary goes up


# ------------------------------------------------------------- add-car form
VALID = dict(make="Lotus", model="Emira", condition="new", price="75,000", power="400",
             top_speed="180", acceleration="4.1", fuel_type="Petrol", transmission="manual",
             gearbox="6 Speed", drivetrain="RWD", capacity="56", weight="1405",
             dimensions="4412X1895x1225", colour="yellow", year="2023", road_tax="190")


def test_valid_form_is_converted_to_typed_fields():
    fields = build_car_fields(VALID)
    assert fields["price"] == 75000 and fields["acceleration"] == 4.1
    assert fields["condition"] == "New" and fields["transmission"] == "Manual"
    assert fields["dimensions"] == "4412 x 1895 x 1225" and fields["colour"] == "Yellow"


def test_optional_fields_accept_na_or_blank():
    fields = build_car_fields({**VALID, "top_speed": "N/A", "weight": "", "dimensions": "n/a"})
    assert fields["top_speed"] is None and fields["weight"] is None
    assert fields["dimensions"] == "N/A"


def test_all_errors_reported_together():
    with pytest.raises(ValidationError) as err:
        build_car_fields({**VALID, "make": " ", "price": "-5", "year": "20x3",
                          "condition": "Broken", "dimensions": "big"})
    assert set(err.value.errors) == {"make", "price", "year", "condition", "dimensions"}


@pytest.mark.parametrize("field, value", [("price", "12.5"), ("power", "0"), ("year", "1850"),
                                          ("acceleration", "fast"), ("weight", "-1")])
def test_individual_bad_values(field, value):
    with pytest.raises(ValidationError) as err:
        build_car_fields({**VALID, field: value})
    assert field in err.value.errors


# ------------------------------------------------------------- storage
def test_old_style_file_is_upgraded(tmp_path):
    """A file shaped like the original treeview data.xlsx (Sheet1, no Stock ID)."""
    path = tmp_path / "old.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Make", "Model", "Used/New", "Price (£)", "Power (bhp)", "Top Speed (mph)",
               "0-60 mph (seconds)", "Fuel Type", "Transmission", "Gearbox", "Drivetrain",
               "Fuel Capacity (L) / Battery Capacity (kWh)", "Weight (kg)",
               "Dimensions Length x Width x Height (mm)", "Colour", "Year", "Annual Road Tax"])
    ws.append(["Toyota", "GT86", "Used", 12250, 197, 140, 7.4, "Petrol", "Manual", "6 Speed",
               "RWD", 50, 1240, "4240 x 1775 x 1285", "black", 2020, 320])
    wb.save(path)

    repo = StockRepository(path)
    assert repo.all()[0].stock_id == 1
    assert openpyxl.load_workbook(path).sheetnames == ["Stock", "Sales"]


def test_missing_file_is_created_empty(tmp_path):
    repo = StockRepository(tmp_path / "new.xlsx")
    assert repo.all() == [] and repo.known_brands() == set()


def test_add_assigns_next_stock_id_and_persists(stock_repo):
    before = len(stock_repo.all())
    car = stock_repo.add(build_car_fields(VALID))
    assert car.stock_id == 37
    reloaded = stock_repo.all()
    assert len(reloaded) == before + 1 and reloaded[-1] == car


def test_sell_moves_car_to_sales_and_keeps_brand_known(stock_repo):
    car = next(c for c in stock_repo.all() if c.make == "Maserati")
    sale = stock_repo.sell(car.stock_id, "MLanner5671", 120_000)

    assert all(c.stock_id != car.stock_id for c in stock_repo.all())
    assert stock_repo.sales() == [sale] and sale.sale_price == 120_000
    assert "Maserati" in stock_repo.known_brands()          # still listed (greyed out in the UI)
    with pytest.raises(StockDataError):
        stock_repo.sell(car.stock_id, "MLanner5671", 1)      # can't sell twice


def test_stock_ids_are_never_reused_after_a_sale(stock_repo):
    last = max(stock_repo.all(), key=lambda c: c.stock_id)
    stock_repo.sell(last.stock_id, "u", 1)
    assert stock_repo.add(build_car_fields(VALID)).stock_id == last.stock_id + 1


def test_missing_column_gives_clear_error(tmp_path):
    path = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Make", "Model"])
    wb.save(path)
    with pytest.raises(StockDataError, match="missing column"):
        StockRepository(path).all()
