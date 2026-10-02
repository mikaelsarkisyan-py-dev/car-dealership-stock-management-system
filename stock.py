"""Cars and sales: the stock.xlsx file, search/filter rules, chart bands and add-car validation."""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.workbook import Workbook

PRICE_BAND = 5_000  # width of each price band in the bar chart (e.g. £10k - £15k)


# ============================================================================
# Data classes
# ============================================================================

NA = "N/A"

SPEC_LABELS = (
    "Power", "Top speed", "0-60 mph", "Fuel type", "Transmission", "Gearbox",
    "Drivetrain", "Fuel (L) / battery (kWh)", "Weight", "Dimensions (L x W x H mm)",
    "Colour", "Year", "Annual road tax", "Price",
)


@dataclass(frozen=True)
class Car:
    stock_id: int
    make: str
    model: str
    condition: str              # "Used" or "New"
    price: int
    power: int                  # bhp
    top_speed: int | None       # mph
    acceleration: float | None  # 0-60 mph, seconds
    fuel_type: str
    transmission: str
    gearbox: str
    drivetrain: str
    capacity: float | None      # fuel (L) or battery (kWh)
    weight: int | None          # kg
    dimensions: str
    colour: str
    year: int
    road_tax: int

    @property
    def name(self) -> str:
        return f"{self.make} {self.model}"

    def spec_values(self) -> list[str]:
        """Display text for each entry in SPEC_LABELS, in the same order."""
        def fmt(value, template="{}"):
            return NA if value is None else template.format(value)

        return [
            fmt(self.power, "{} bhp"),
            fmt(self.top_speed, "{} mph"),
            fmt(self.acceleration, "{} secs"),
            self.fuel_type,
            self.transmission,
            self.gearbox,
            self.drivetrain,
            fmt(self.capacity, "{:g}"),
            fmt(self.weight, "{} kg"),
            self.dimensions,
            self.colour,
            str(self.year),
            f"£{self.road_tax:,}",
            f"£{self.price:,}",
        ]

    def spec_rows(self) -> list[tuple[str, str]]:
        """(label, value) pairs for the specs views."""
        return list(zip(SPEC_LABELS, self.spec_values()))


@dataclass(frozen=True)
class Sale:
    stock_id: int
    make: str
    model: str
    date: datetime
    user_id: str
    sale_price: int


# ============================================================================
# Excel storage (Stock and Sales sheets)
# ============================================================================

STOCK_SHEET = "Stock"
SALES_SHEET = "Sales"

# Car attribute -> column header in the workbook (order = column order for new files)
STOCK_HEADERS = {
    "make": "Make",
    "model": "Model",
    "condition": "Used/New",
    "price": "Price (£)",
    "power": "Power (bhp)",
    "top_speed": "Top Speed (mph)",
    "acceleration": "0-60 mph (seconds)",
    "fuel_type": "Fuel Type",
    "transmission": "Transmission",
    "gearbox": "Gearbox",
    "drivetrain": "Drivetrain",
    "capacity": "Fuel Capacity (L) / Battery Capacity (kWh)",
    "weight": "Weight (kg)",
    "dimensions": "Dimensions Length x Width x Height (mm)",
    "colour": "Colour",
    "year": "Year",
    "road_tax": "Annual Road Tax",
    "stock_id": "Stock ID",
}
SALES_HEADERS = ["Stock ID", "Make", "Model", "Date", "User ID", "Sale Price (£)"]


class StockDataError(Exception):
    """Raised for unreadable/unwritable stock data; message is safe to show users."""


# ---------------------------------------------------------------- cell parsing
def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and value.strip().upper() in ("", NA))


def _opt_int(value) -> int | None:
    return None if _is_blank(value) else int(float(value))


def _opt_float(value) -> float | None:
    return None if _is_blank(value) else float(value)


def _text(value) -> str:
    return "" if value is None else str(value).strip()


class StockRepository:
    def __init__(self, path: Path):
        self.path = Path(path)
        self._prepare_file()

    # ------------------------------------------------------------ file plumbing
    def _prepare_file(self) -> None:
        """Create the workbook if missing, or upgrade an older layout in place."""
        if not self.path.exists():
            workbook = Workbook()
            workbook.active.title = STOCK_SHEET
            workbook[STOCK_SHEET].append(list(STOCK_HEADERS.values()))
            workbook.create_sheet(SALES_SHEET).append(SALES_HEADERS)
            self._save(workbook)
            return

        workbook = self._open()
        changed = False

        sheet = self._stock_sheet(workbook)
        if sheet.title != STOCK_SHEET:
            sheet.title = STOCK_SHEET
            changed = True

        headers = self._header_map(sheet)
        if STOCK_HEADERS["stock_id"] not in headers:
            column = sheet.max_column + 1
            sheet.cell(row=1, column=column, value=STOCK_HEADERS["stock_id"])
            number = 0
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row=row, column=1).value is not None:
                    number += 1
                    sheet.cell(row=row, column=column, value=number)
            changed = True

        if SALES_SHEET not in workbook.sheetnames:
            workbook.create_sheet(SALES_SHEET).append(SALES_HEADERS)
            changed = True

        if changed:
            self._save(workbook)

    def _open(self) -> Workbook:
        try:
            return openpyxl.load_workbook(self.path)
        except FileNotFoundError as exc:
            raise StockDataError(f"Stock file not found: {self.path}") from exc
        except PermissionError as exc:
            raise StockDataError("Cannot open the stock file. Is it open in Excel?") from exc

    def _save(self, workbook: Workbook) -> None:
        try:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(self.path)
        except PermissionError as exc:
            raise StockDataError("Cannot save the stock file. Close it in Excel and try again.") from exc

    @staticmethod
    def _stock_sheet(workbook: Workbook):
        return workbook[STOCK_SHEET] if STOCK_SHEET in workbook.sheetnames else workbook.worksheets[0]

    @staticmethod
    def _header_map(sheet) -> dict[str, int]:
        """header text -> 1-based column index"""
        return {
            str(cell.value).strip(): cell.column
            for cell in sheet[1]
            if cell.value is not None
        }

    # ------------------------------------------------------------------ reading
    def all(self) -> list[Car]:
        sheet = self._stock_sheet(self._open())
        headers = self._header_map(sheet)
        missing = [h for h in STOCK_HEADERS.values() if h not in headers]
        if missing:
            raise StockDataError(f"Stock file is missing column(s): {', '.join(missing)}")

        def cell(row, attr):
            return row[headers[STOCK_HEADERS[attr]] - 1]

        cars = []
        for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if _is_blank(cell(row, "make")) and _is_blank(cell(row, "model")):
                continue  # empty row
            try:
                cars.append(Car(
                    stock_id=int(cell(row, "stock_id")),
                    make=_text(cell(row, "make")),
                    model=_text(cell(row, "model")),
                    condition=_text(cell(row, "condition")).title(),
                    price=int(float(cell(row, "price"))),
                    power=int(float(cell(row, "power"))),
                    top_speed=_opt_int(cell(row, "top_speed")),
                    acceleration=_opt_float(cell(row, "acceleration")),
                    fuel_type=_text(cell(row, "fuel_type")),
                    transmission=_text(cell(row, "transmission")),
                    gearbox=_text(cell(row, "gearbox")),
                    drivetrain=_text(cell(row, "drivetrain")),
                    capacity=_opt_float(cell(row, "capacity")),
                    weight=_opt_int(cell(row, "weight")),
                    dimensions=_text(cell(row, "dimensions")) or NA,
                    colour=_text(cell(row, "colour")).title(),
                    year=int(float(cell(row, "year"))),
                    road_tax=int(float(cell(row, "road_tax"))),
                ))
            except (TypeError, ValueError) as exc:
                raise StockDataError(f"Row {number} of the stock file has invalid data.") from exc
        return cars

    def known_brands(self) -> set[str]:
        """Every brand ever stocked: current stock plus anything already sold.

        This lets the UI grey out a brand once its last car is sold.
        """
        brands = {car.make for car in self.all()}
        for sale in self.sales():
            brands.add(sale.make)
        return brands

    def sales(self) -> list[Sale]:
        sheet = self._open()[SALES_SHEET]
        sales = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            stock_id, make, model, date, user_id, price = row[:6]
            sales.append(Sale(int(stock_id), str(make), str(model), date, str(user_id), int(price)))
        return sales

    # ------------------------------------------------------------------ writing
    def add(self, fields: dict) -> Car:
        """Append a car. `fields` holds every Car attribute except stock_id."""
        workbook = self._open()
        sheet = self._stock_sheet(workbook)
        headers = self._header_map(sheet)

        id_column = headers[STOCK_HEADERS["stock_id"]]
        existing_ids = [
            sheet.cell(row=r, column=id_column).value for r in range(2, sheet.max_row + 1)
        ]
        sold_ids = [sale.stock_id for sale in self.sales()]
        # Never reuse an ID, even one belonging to a car that has since been sold.
        stock_id = max([int(i) for i in existing_ids if i is not None] + sold_ids, default=0) + 1
        values = {**fields, "stock_id": stock_id}

        new_row = sheet.max_row + 1
        for attr, header in STOCK_HEADERS.items():
            value = values[attr]
            sheet.cell(row=new_row, column=headers[header], value=NA if value is None else value)
        self._save(workbook)
        return Car(**values)

    def sell(self, stock_id: int, user_id: str, sale_price: int) -> Sale:
        """Remove a car from stock and record the sale."""
        workbook = self._open()
        sheet = self._stock_sheet(workbook)
        headers = self._header_map(sheet)
        id_column = headers[STOCK_HEADERS["stock_id"]]

        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row=row, column=id_column).value == stock_id:
                make = _text(sheet.cell(row=row, column=headers[STOCK_HEADERS["make"]]).value)
                model = _text(sheet.cell(row=row, column=headers[STOCK_HEADERS["model"]]).value)
                sheet.delete_rows(row)
                sale = Sale(stock_id, make, model, datetime.now().replace(microsecond=0),
                            user_id, sale_price)
                workbook[SALES_SHEET].append(
                    [sale.stock_id, sale.make, sale.model, sale.date, sale.user_id, sale.sale_price]
                )
                self._save(workbook)
                return sale
        raise StockDataError(f"Stock ID {stock_id} was not found (already sold?).")


# ============================================================================
# Filtering, grouping, chart bands, add-car validation
# ============================================================================

class ValidationError(Exception):
    """`errors` maps a form field name to a message."""

    def __init__(self, errors: dict[str, str]):
        self.errors = errors
        super().__init__("\n".join(errors.values()))


# ------------------------------------------------------------------ filtering
def _same(a: str, b: str) -> bool:
    return a.casefold() == b.casefold()


@dataclass(frozen=True)
class CarFilter:
    """All set criteria must match (AND). Unset (None/empty) criteria are ignored."""
    text: str = ""
    make: str | None = None
    model: str | None = None
    condition: str | None = None
    transmission: str | None = None
    fuel_type: str | None = None
    gearbox: str | None = None
    drivetrain: str | None = None
    colour: str | None = None
    min_price: int | None = None
    max_price: int | None = None

    def matches(self, car: Car) -> bool:
        if self.text and self.text.strip().casefold() not in car.name.casefold():
            return False
        for attr in ("make", "model", "condition", "transmission",
                     "fuel_type", "gearbox", "drivetrain", "colour"):
            wanted = getattr(self, attr)
            if wanted and not _same(wanted, getattr(car, attr)):
                return False
        if self.min_price is not None and car.price < self.min_price:
            return False
        if self.max_price is not None and car.price > self.max_price:
            return False
        return True


def filter_cars(cars: list[Car], criteria: CarFilter) -> list[Car]:
    return [car for car in cars if criteria.matches(car)]


def distinct_values(cars: list[Car], attr: str) -> list[str]:
    """Unique values of an attribute, case-insensitive, sorted for display."""
    seen: dict[str, str] = {}
    for car in cars:
        value = getattr(car, attr)
        if value:
            seen.setdefault(value.casefold(), value)
    return sorted(seen.values(), key=str.casefold)


# ------------------------------------------------------------------- grouping
def brand_counts(cars: list[Car], known_brands: set[str] = frozenset()) -> dict[str, int]:
    """{brand: cars in stock}, including brands that are known but now at 0."""
    counts = {brand: 0 for brand in known_brands}
    for car in cars:
        counts[car.make] = counts.get(car.make, 0) + 1
    return dict(sorted(counts.items(), key=lambda item: item[0].casefold()))


def cars_for(cars: list[Car], make: str, model: str | None = None) -> list[Car]:
    return [c for c in cars
            if _same(c.make, make) and (model is None or _same(c.model, model))]


def model_counts(cars: list[Car], make: str) -> dict[str, int]:
    counts: dict[str, int] = {}
    for car in cars_for(cars, make):
        counts[car.model] = counts.get(car.model, 0) + 1
    return dict(sorted(counts.items(), key=lambda item: item[0].casefold()))


# -------------------------------------------------------------- price bands
def _thousands(amount: int) -> str:
    return f"£{amount / 1000:g}k"


def price_distribution(cars: list[Car], band: int = PRICE_BAND) -> list[tuple[str, int]]:
    """[(label, count)] for consecutive price bands, e.g. ("£10k - £15k", 2)."""
    if not cars:
        return []
    first = min(c.price for c in cars) // band
    last = max(c.price for c in cars) // band
    counts = {i: 0 for i in range(first, last + 1)}
    for car in cars:
        counts[car.price // band] += 1
    return [(f"{_thousands(i * band)} - {_thousands((i + 1) * band)}", n)
            for i, n in counts.items()]


# ------------------------------------------------------- add-car validation
_DIMENSIONS_RE = re.compile(r"^(\d{3,5})\s*x\s*(\d{3,5})\s*x\s*(\d{3,5}|N/A)$", re.IGNORECASE)
_TEXT_FIELDS = {
    "make": "Make", "model": "Model", "fuel_type": "Fuel type", "gearbox": "Gearbox",
    "drivetrain": "Drivetrain", "colour": "Colour",
}
_CHOICES = {
    "condition": ("Used / New", ("Used", "New")),
    "transmission": ("Transmission", ("Manual", "Automatic")),
}
_REQUIRED_NUMBERS = {  # attr: (label, cast, minimum)
    "price": ("Price", int, 1),
    "power": ("Power", int, 1),
    "road_tax": ("Annual road tax", int, 0),
}
_OPTIONAL_NUMBERS = {  # may be left as N/A
    "top_speed": ("Top speed", int),
    "acceleration": ("0-60 time", float),
    "capacity": ("Fuel / battery capacity", float),
    "weight": ("Weight", int),
}


def _number(text: str, cast):
    value = float(text.replace(",", "").replace("£", "").strip())
    if cast is int and value != int(value):
        raise ValueError
    return cast(value)


def build_car_fields(form: dict[str, str]) -> dict:
    """Validate raw form text. Returns typed Car fields (minus stock_id) or raises
    ValidationError listing every problem at once."""
    errors: dict[str, str] = {}
    fields: dict = {}

    for attr, label in _TEXT_FIELDS.items():
        value = form.get(attr, "").strip()
        if value:
            fields[attr] = value
        else:
            errors[attr] = f"{label} is required."

    for attr, (label, options) in _CHOICES.items():
        value = form.get(attr, "").strip()
        match = next((o for o in options if o.casefold() == value.casefold()), None)
        if match:
            fields[attr] = match
        else:
            errors[attr] = f"{label}: choose {' or '.join(options)}."

    for attr, (label, cast, minimum) in _REQUIRED_NUMBERS.items():
        try:
            value = _number(form.get(attr, ""), cast)
            if value < minimum:
                raise ValueError
            fields[attr] = value
        except ValueError:
            errors[attr] = f"{label} must be a whole number{' above 0' if minimum else ''}."

    for attr, (label, cast) in _OPTIONAL_NUMBERS.items():
        text = form.get(attr, "").strip()
        if text.upper() in ("", NA):
            fields[attr] = None
            continue
        try:
            value = _number(text, cast)
            if value <= 0:
                raise ValueError
            fields[attr] = value
        except ValueError:
            errors[attr] = f"{label} must be a positive number (or N/A)."

    try:
        year = int(form.get("year", "").strip())
        if not 1900 <= year <= date.today().year + 1:
            raise ValueError
        fields["year"] = year
    except ValueError:
        errors["year"] = f"Year must be a 4-digit year (1900-{date.today().year + 1})."

    dimensions = form.get("dimensions", "").strip()
    match = _DIMENSIONS_RE.match(dimensions)
    if dimensions.upper() == NA:
        fields["dimensions"] = NA
    elif match:
        length, width, height = (part.upper() for part in match.groups())
        fields["dimensions"] = f"{length} x {width} x {height}"
    else:
        errors["dimensions"] = "Dimensions must look like 4240 x 1775 x 1285 (or N/A)."

    if errors:
        raise ValidationError(errors)
    fields["colour"] = fields["colour"].title()
    return fields
